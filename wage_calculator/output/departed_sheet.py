"""5.7절: 당월 A파일에는 없지만(완전퇴사) 소급조정이 발생한 사람의 내역을
당월 계산 결과와 별도 워크북으로 출력하는 시트 2종."""
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SHEET_NAME = "소급대상자 내역"
EVIDENCE_SHEET_NAME = "산정근거(수식)"
DATA_START_ROW = 2

COL = {
    "seq": 1, "name": 2, "ssn": 3, "bank": 4, "account": 5,
    "prev_recalculated": 6, "prev_paid": 7, "adjustment": 8,
}

HEADERS = {
    "seq": "순번", "name": "성명", "ssn": "주민번호", "bank": "거래은행", "account": "계좌번호",
    "prev_recalculated": "전월재계산액", "prev_paid": "전월실지급액",
    "adjustment": "소급조정액(최종지급액)",
}

EVIDENCE_COL = {"seq": 1, "name": 2, "prev_recalculated": 3, "prev_paid": 4, "adjustment": 5}
EVIDENCE_HEADERS = {
    "seq": "순번", "name": "성명", "prev_recalculated": "전월재계산액",
    "prev_paid": "전월실지급액", "adjustment": "소급조정액(=전월재계산액-전월실지급액)",
}


def build_departed_sheet(wb, departed_results):
    """departed_results: list[core.retroactive.DepartedRetro]"""
    ws = wb.create_sheet(SHEET_NAME)
    for key, col in COL.items():
        cell = ws.cell(row=1, column=col, value=HEADERS[key])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 16

    row = DATA_START_ROW
    for i, d in enumerate(departed_results, start=1):
        ws.cell(row=row, column=COL["seq"], value=i)
        ws.cell(row=row, column=COL["name"], value=d.name)
        ws.cell(row=row, column=COL["ssn"], value=d.ssn)
        ws.cell(row=row, column=COL["bank"], value=d.bank)
        ws.cell(row=row, column=COL["account"], value=d.account)
        ws.cell(row=row, column=COL["prev_recalculated"], value=d.prev_recalculated)
        ws.cell(row=row, column=COL["prev_paid"], value=d.prev_paid)
        ws.cell(row=row, column=COL["adjustment"], value=d.adjustment)
        row += 1

    for col in (COL["prev_recalculated"], COL["prev_paid"], COL["adjustment"]):
        for r_ in range(DATA_START_ROW, row):
            ws.cell(row=r_, column=col).number_format = "#,##0"

    return ws


def build_departed_evidence_sheet(wb, departed_results):
    """§6.3: 소급조정액=최종지급액이므로 지급총액 관련 행은 생략하고
    소급조정액 수식(=전월재계산액-전월실지급액)만 표시하는 산정근거 시트."""
    ws = wb.create_sheet(EVIDENCE_SHEET_NAME)
    for key, col in EVIDENCE_COL.items():
        cell = ws.cell(row=1, column=col, value=EVIDENCE_HEADERS[key])
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = 16

    row = DATA_START_ROW
    for i, d in enumerate(departed_results, start=1):
        ws.cell(row=row, column=EVIDENCE_COL["seq"], value=i)
        ws.cell(row=row, column=EVIDENCE_COL["name"], value=d.name)
        ws.cell(row=row, column=EVIDENCE_COL["prev_recalculated"], value=d.prev_recalculated)
        ws.cell(row=row, column=EVIDENCE_COL["prev_paid"], value=d.prev_paid)
        recalc_addr = f"{get_column_letter(EVIDENCE_COL['prev_recalculated'])}{row}"
        paid_addr = f"{get_column_letter(EVIDENCE_COL['prev_paid'])}{row}"
        ws.cell(row=row, column=EVIDENCE_COL["adjustment"], value=f"={recalc_addr}-{paid_addr}")
        row += 1

    for col in (EVIDENCE_COL["prev_recalculated"], EVIDENCE_COL["prev_paid"], EVIDENCE_COL["adjustment"]):
        for r_ in range(DATA_START_ROW, row):
            ws.cell(row=r_, column=col).number_format = "#,##0"

    return ws
