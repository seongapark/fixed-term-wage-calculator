"""B파일(근무상황) 원본 행을 그대로 옮긴 시트.

이 시트는 단순한 보관용이 아니라 **다음 달 계산의 입력**이다. 다음 달에 이 파일을
'전월 임금내역'으로 첨부하면 core/parser.py가 이 시트를 읽어 지난 달 근무상황을
복원한다(계약 시작월부터 매번 다시 조회해 첨부하지 않아도 되는 이유). 그래서 열
구성을 바꿀 때는 core/parser.load_previous_status_rows()와 같이 고쳐야 한다.

'지급판정'·'주휴연가판정'은 확인 화면에서 사람이 고른 결과다. 규칙만으로 정해진
행은 비워 둔다 - 다음 달에 규칙으로 다시 판정해도 같은 값이 나오므로 저장할 이유가
없고, 비어 있으면 규칙이 바뀌었을 때 새 규칙이 적용된다.
"""
from openpyxl.styles import Font

from core import pending

SHEET_NAME = "근무상황(기간제)"

COLUMNS = [
    "소속", "직급", "성명", "생년월일", "종별", "사용기간(날짜)", "사용시간(시분)",
    "사유", "연락처", "결재상태", "비고", pending.PAID_COL, pending.ACCRUAL_COL,
]


def build_raw_status_sheet(wb, giganje_rows):
    ws = wb.create_sheet(SHEET_NAME)
    for c, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = Font(bold=True)

    for r, row in enumerate(giganje_rows, start=2):
        for c, header in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=c, value=row.get(header))

    for c in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 16

    return ws
