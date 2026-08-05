"""7-0: '임금내역(월중)' 시트 재현(열 구성/헤더/순서를 원본과 동일하게, 값은 계산 결과)."""
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

SHEET_NAME = "임금내역(월중)"

HEADER_ROWS = 3
DATA_START_ROW = HEADER_ROWS + 1

# 컬럼 순서(원본 '임금내역(월중)' 시트와 동일): 1~31열
COL = {
    "seq": 1, "name": 2, "survey": 3, "period_start": 4, "period_end": 5,
    "daily_wage": 6, "actual_workdays": 7, "public_leave": 8, "paid_holiday": 9,
    "absence": 10, "total_days": 11, "late_out": 12, "weekly_holiday": 13,
    "calendar_days": 14, "meal_eligible": 15, "remaining_leave": 16,
    "gross_pay": 17, "late_out_deduction": 18, "base_pay": 19, "weekly_holiday_pay": 20,
    "meal_allowance": 21, "leave_compensation": 22, "total_payment": 23,
    "ssn": 24, "contract_start": 25, "contract_end": 26, "bank": 27, "account": 28,
    "note": 29, "retro_adjustment": 30, "final_payment": 31,
}


def _set(ws, addr, value, merge=None):
    ws[addr] = value
    ws[addr].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws[addr].font = Font(bold=True)
    if merge:
        ws.merge_cells(merge)


def _write_headers(ws):
    _set(ws, "A1", "순번", "A1:A3")
    _set(ws, "B1", "성  명", "B1:B3")
    _set(ws, "C1", "조사", "C1:C3")
    _set(ws, "D1", "급여계산 기간", "D1:E2")
    _set(ws, "D3", "시작")
    _set(ws, "E3", "마지막")
    _set(ws, "F1", "일  급", "F1:F3")

    _set(ws, "G1", "산정내역", "G1:P1")
    _set(ws, "G3", "실출근\n(일)")
    _set(ws, "H3", "공가\n(일)")
    _set(ws, "I3", "유급\n휴일\n(일)")
    _set(ws, "J3", "(결근)")
    _set(ws, "K2", "계\n(일)", "K2:K3")
    _set(ws, "L2", "조퇴\n외출\n(시간)", "L2:L3")
    _set(ws, "M2", "주휴\n(일)", "M2:M3")
    _set(ws, "N3", "월력상")
    _set(ws, "O2", "식대\n해당일", "O2:O3")
    _set(ws, "P2", "잔여\n연가\n(일)", "P2:P3")

    _set(ws, "Q1", "지급내역", "Q1:V1")
    _set(ws, "Q3", "급여액")
    _set(ws, "R3", "조퇴외출\n공제")
    _set(ws, "S2", "기본급\n(급여-공제)", "S2:S3")
    _set(ws, "T2", "주휴수당", "T2:T3")
    _set(ws, "U2", "정액급식비", "U2:U3")
    _set(ws, "V2", "연가보상비", "V2:V3")

    _set(ws, "W1", "기존지급액", "W1:W3")
    _set(ws, "X1", "주민번호", "X1:X3")
    _set(ws, "Y1", "계약일자", "Y1:Z2")
    _set(ws, "Y3", "시작")
    _set(ws, "Z3", "마지막")
    _set(ws, "AA1", "거래\n은행", "AA1:AA3")
    _set(ws, "AB1", "계 좌 번  호", "AB1:AB3")
    _set(ws, "AC1", "비고", "AC1:AC3")
    _set(ws, "AD1", "소급\n조정액", "AD1:AD3")
    _set(ws, "AE1", "최종\n지급액", "AE1:AE3")

    for col in range(1, 32):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = 10

    ws["H3"].comment = Comment(
        "공가(일)에는 유급 특별휴가 일수가 합산되어 있습니다. 세부 날짜는 비고란 참고.",
        "임금계산 프로그램",
    )
    ws["J3"].comment = Comment(
        "(결근)에는 무급 특별휴가 일수가 합산되어 있습니다. 세부 날짜는 비고란 참고.",
        "임금계산 프로그램",
    )


def _num(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def build_wage_sheet(wb, results, seq_start=1, retro_adjustments=None):
    from core.parser import person_key

    retro_adjustments = retro_adjustments or {}
    ws = wb.create_sheet(SHEET_NAME)
    _write_headers(ws)

    row = DATA_START_ROW
    for i, r in enumerate(results, start=seq_start):
        ws.cell(row=row, column=COL["seq"], value=i)
        ws.cell(row=row, column=COL["name"], value=r.name)
        ws.cell(row=row, column=COL["survey"], value=r.survey_name)
        ws.cell(row=row, column=COL["period_start"], value=r.period_start)
        ws.cell(row=row, column=COL["period_end"], value=r.period_end)
        ws.cell(row=row, column=COL["daily_wage"], value=r.daily_wage)
        ws.cell(row=row, column=COL["actual_workdays"], value=_num(r.actual_workdays))
        ws.cell(row=row, column=COL["public_leave"], value=_num(r.public_leave_days))
        ws.cell(row=row, column=COL["paid_holiday"], value=r.paid_holiday_days)
        ws.cell(row=row, column=COL["absence"], value=_num(r.absence_days))
        ws.cell(row=row, column=COL["total_days"], value=_num(r.total_days))
        ws.cell(row=row, column=COL["late_out"], value=_num(round(r.late_out_minutes / 60, 2)))
        ws.cell(row=row, column=COL["weekly_holiday"], value=r.weekly_holiday_days)
        ws.cell(row=row, column=COL["calendar_days"], value=r.calendar_month_days)
        ws.cell(row=row, column=COL["meal_eligible"], value=_num(r.meal_eligible_days))
        ws.cell(row=row, column=COL["remaining_leave"], value=round(r.remaining_leave_days, 4))
        ws.cell(row=row, column=COL["gross_pay"], value=_num(r.gross_pay))
        ws.cell(row=row, column=COL["late_out_deduction"], value=r.late_out_deduction)
        ws.cell(row=row, column=COL["base_pay"], value=r.base_pay)
        ws.cell(row=row, column=COL["weekly_holiday_pay"], value=r.weekly_holiday_pay)
        ws.cell(row=row, column=COL["meal_allowance"], value=r.meal_allowance)
        ws.cell(row=row, column=COL["leave_compensation"], value=r.leave_compensation)
        ws.cell(row=row, column=COL["total_payment"], value=r.total_payment)
        ws.cell(row=row, column=COL["ssn"], value=r.ssn)
        ws.cell(row=row, column=COL["contract_start"], value=r.contract_start)
        ws.cell(row=row, column=COL["contract_end"], value=r.contract_end)
        ws.cell(row=row, column=COL["bank"], value=r.bank)
        ws.cell(row=row, column=COL["account"], value=r.account)
        ws.cell(row=row, column=COL["note"], value=r.special_leave_note)
        adjustment = retro_adjustments.get(person_key(r.name, r.birth), 0)
        ws.cell(row=row, column=COL["retro_adjustment"], value=adjustment)
        ws.cell(row=row, column=COL["final_payment"], value=r.total_payment + adjustment)
        row += 1

    for col in (COL["period_start"], COL["period_end"], COL["contract_start"], COL["contract_end"]):
        for r_ in range(DATA_START_ROW, row):
            ws.cell(row=r_, column=col).number_format = "yyyy-mm-dd"

    return ws
