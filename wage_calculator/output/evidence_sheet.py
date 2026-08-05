"""6장(지급내역 수식 근거): 지급내역(금액) 계산 과정을 실제 엑셀 수식으로
재현하는 시트. 절삭 지점(4곳: 조퇴외출공제/기본급/정액급식비/지급총액,
모두 10원 단위 - core/payroll.py 참고)은 바꾸지 않고, 이미 계산된 값을
사람이 엑셀에서 셀 단위로 추적 검증할 수 있게 하는 것이 목적이다.
실출근일수·주휴·연차 발생 여부 같은 판정 로직(core/leave_engine.py)은
분기가 많아 수식으로 옮기지 않고 입력값 그대로 보여준다."""
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SHEET_NAME = "산정근거(수식)"
HEADER_ROW = 1
DATA_START_ROW = 2

COL = {
    "seq": 1, "name": 2,
    "hourly_wage": 3, "total_days": 4, "late_out_minutes": 5, "weekly_holiday_days": 6,
    "meal_allowance_rate": 7, "calendar_month_days": 8, "meal_eligible_days": 9,
    "remaining_leave_days": 10, "is_final_month": 11,
    "daily_wage": 12, "daily_meal": 13, "gross_pay": 14, "late_out_deduction": 15,
    "base_pay": 16, "weekly_holiday_pay": 17, "meal_allowance": 18, "leave_compensation": 19,
    "total_payment": 20, "retro_adjustment": 21, "final_payment": 22,
}

HEADERS = {
    "seq": "순번", "name": "성명", "hourly_wage": "시급", "total_days": "계(일)",
    "late_out_minutes": "조퇴외출(분)", "weekly_holiday_days": "주휴(일)",
    "meal_allowance_rate": "월식대", "calendar_month_days": "월력상",
    "meal_eligible_days": "식대해당일", "remaining_leave_days": "잔여연가(일)",
    "is_final_month": "최종월여부", "daily_wage": "일급", "daily_meal": "일급식대",
    "gross_pay": "급여액", "late_out_deduction": "조퇴외출공제", "base_pay": "기본급",
    "weekly_holiday_pay": "주휴수당", "meal_allowance": "정액급식비",
    "leave_compensation": "연가보상비", "total_payment": "지급총액",
    "retro_adjustment": "소급조정액", "final_payment": "최종지급액",
}


def _addr(row, key):
    return f"{get_column_letter(COL[key])}{row}"


def _write_headers(ws):
    for key, col in COL.items():
        cell = ws.cell(row=HEADER_ROW, column=col, value=HEADERS[key])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 12


def build_evidence_sheet(wb, results, config, retro_adjustments=None):
    """results: list[PayrollResult]. retro_adjustments: dict[person_key, int]."""
    from core.parser import person_key

    retro_adjustments = retro_adjustments or {}
    ws = wb.create_sheet(SHEET_NAME)
    _write_headers(ws)

    row = DATA_START_ROW
    for i, r in enumerate(results, start=1):
        ws.cell(row=row, column=COL["seq"], value=i)
        ws.cell(row=row, column=COL["name"], value=r.name)

        # 입력값(고정값): 이미 파이썬에서 계산된 값을 그대로 적어 넣는다.
        ws.cell(row=row, column=COL["hourly_wage"], value=config.hourly_wage_for(r.period_start.year))
        ws.cell(row=row, column=COL["total_days"], value=r.total_days)
        ws.cell(row=row, column=COL["late_out_minutes"], value=r.late_out_minutes)
        ws.cell(row=row, column=COL["weekly_holiday_days"], value=r.weekly_holiday_days)
        ws.cell(row=row, column=COL["meal_allowance_rate"], value=config.meal_allowance_for(r.period_start.year))
        ws.cell(row=row, column=COL["calendar_month_days"], value=r.calendar_month_days)
        ws.cell(row=row, column=COL["meal_eligible_days"], value=r.meal_eligible_days)
        ws.cell(row=row, column=COL["remaining_leave_days"], value=r.remaining_leave_days)
        ws.cell(row=row, column=COL["is_final_month"], value=1 if r.is_final_month else 0)

        # 수식 셀: 같은 행의 입력값 셀을 참조하는 실제 엑셀 ROUNDDOWN 수식.
        ws.cell(row=row, column=COL["daily_wage"],
                value=f"=ROUNDDOWN({_addr(row,'hourly_wage')}*8,0)")
        ws.cell(row=row, column=COL["daily_meal"],
                value=f"=ROUNDDOWN({_addr(row,'meal_allowance_rate')}/209*8,0)")
        ws.cell(row=row, column=COL["gross_pay"],
                value=f"={_addr(row,'daily_wage')}*{_addr(row,'total_days')}")
        ws.cell(row=row, column=COL["late_out_deduction"],
                value=f"=ROUNDDOWN({_addr(row,'hourly_wage')}/60*{_addr(row,'late_out_minutes')},-1)")
        ws.cell(row=row, column=COL["base_pay"],
                value=f"=ROUNDDOWN({_addr(row,'gross_pay')}-{_addr(row,'late_out_deduction')},-1)")
        ws.cell(row=row, column=COL["weekly_holiday_pay"],
                value=f"={_addr(row,'daily_wage')}*{_addr(row,'weekly_holiday_days')}")
        ws.cell(row=row, column=COL["meal_allowance"],
                value=f"=ROUNDDOWN({_addr(row,'meal_allowance_rate')}/{_addr(row,'calendar_month_days')}*{_addr(row,'meal_eligible_days')},-1)")
        ws.cell(row=row, column=COL["leave_compensation"],
                value=(
                    f"=IF({_addr(row,'is_final_month')}=1,"
                    f"ROUNDDOWN(({_addr(row,'daily_wage')}+{_addr(row,'daily_meal')})*{_addr(row,'remaining_leave_days')},0),0)"
                ))
        ws.cell(row=row, column=COL["total_payment"],
                value=(
                    f"=ROUNDDOWN({_addr(row,'base_pay')}+{_addr(row,'weekly_holiday_pay')}"
                    f"+{_addr(row,'meal_allowance')}+{_addr(row,'leave_compensation')},-1)"
                ))

        adjustment = retro_adjustments.get(person_key(r.name, r.birth), 0)
        ws.cell(row=row, column=COL["retro_adjustment"], value=adjustment)
        ws.cell(row=row, column=COL["final_payment"],
                value=f"={_addr(row,'total_payment')}+{_addr(row,'retro_adjustment')}")

        row += 1
    return ws
