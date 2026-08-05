import sys
from pathlib import Path
from datetime import date

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from core.parser import person_key
from core.payroll import PayrollResult
from output.wage_sheet import build_wage_sheet


def _result(name, birth, total_payment):
    return PayrollResult(
        name=name, birth=birth, ssn="", bank="", account="", survey_name="테스트조사",
        period_start=date(2026, 7, 1), period_end=date(2026, 7, 31),
        contract_start=date(2026, 7, 1), contract_end=date(2026, 7, 31),
        daily_wage=78560, actual_workdays=23, public_leave_days=0, paid_holiday_days=0,
        absence_days=0, total_days=23, late_out_minutes=0, weekly_holiday_days=4,
        calendar_month_days=31, meal_eligible_days=31, remaining_leave_days=0.0,
        is_final_month=False, gross_pay=1806880, late_out_deduction=0, base_pay=1806880,
        weekly_holiday_pay=314240, meal_allowance=160000, leave_compensation=0,
        total_payment=total_payment,
    )


def test_retro_columns_written_and_summed():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    result = _result("김철수", "19900101", 2281120)
    key = person_key("김철수", "19900101")
    ws = build_wage_sheet(wb, [result], retro_adjustments={key: -10000})

    assert ws.cell(row=4, column=30).value == -10000, "소급조정액(AD열)"
    assert ws.cell(row=4, column=31).value == 2271120, "최종지급액(AE열) = 지급총액+소급조정액"
    print("OK: test_retro_columns_written_and_summed")


def test_no_retro_adjustment_defaults_to_zero():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    result = _result("박신입", "19950505", 1500000)
    ws = build_wage_sheet(wb, [result])

    assert ws.cell(row=4, column=30).value == 0
    assert ws.cell(row=4, column=31).value == 1500000
    print("OK: test_no_retro_adjustment_defaults_to_zero")


if __name__ == "__main__":
    test_retro_columns_written_and_summed()
    test_no_retro_adjustment_defaults_to_zero()
    print("ALL OK")
