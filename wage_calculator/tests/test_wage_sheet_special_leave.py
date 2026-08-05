import sys
from pathlib import Path
from datetime import date
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from core.config import Config
from core.payroll import calc_payroll
from core.models import LeaveEvent
from output.wage_sheet import build_wage_sheet


def test_note_column_and_header_comments():
    p1, p2 = date(2026, 7, 20), date(2026, 7, 21)
    u1 = date(2026, 7, 27)
    events = [
        LeaveEvent(raw_category="특별휴가", classified="유급특별휴가", d=d, is_time_based=False,
                   day_weight=1.0, breaks=False, source_range=(p1, p2))
        for d in (p1, p2)
    ]
    events.append(
        LeaveEvent(raw_category="특별휴가", classified="무급특별휴가", d=u1, is_time_based=False,
                   day_weight=1.0, breaks=False, source_range=(u1, u1))
    )
    person = SimpleNamespace(
        contract_start=date(2026, 7, 1), contract_end=date(2026, 7, 31), events=events,
        name="테스트", birth="", ssn="", bank="", account="", survey_name="테스트조사",
    )
    config = Config({"surveys": [], "rates": {"2026": {"hourly_wage": 9820, "meal_allowance": 160000}}, "holidays": []})
    result = calc_payroll(person, config, 2026, 7)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = build_wage_sheet(wb, [result])

    assert ws["AC4"].value == "특별휴가(유급) 7/20~7/21, 특별휴가(무급) 7/27", ws["AC4"].value
    assert ws["H3"].comment is not None, "H3(공가)에 코멘트가 없음"
    assert "유급 특별휴가" in ws["H3"].comment.text, ws["H3"].comment.text
    assert ws["J3"].comment is not None, "J3(결근)에 코멘트가 없음"
    assert "무급 특별휴가" in ws["J3"].comment.text, ws["J3"].comment.text
    print("OK: test_note_column_and_header_comments")


if __name__ == "__main__":
    test_note_column_and_header_comments()
    print("ALL OK")
