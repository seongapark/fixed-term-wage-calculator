import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from core import pending
from core.config import Config
from output.raw_sheet import SHEET_NAME as STATUS_SHEET_NAME
from webapp.state import AppState


def _config():
    return Config({
        "surveys": [{"name": "테스트조사", "start": "2026-07-01", "end": "2026-09-30"}],
        "rates": {"2026": {"daily_wage": 78560, "meal_allowance": 160000}},
        "holidays": [],
    })


def _write_a(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["성명", "주민번호", "은행", "계좌번호", "생년월일"])
    ws.append(["홍길동", "900101-1234567", "농협", "123-456", "1990-01-01"])
    wb.save(path)


def _write_b(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["소속", "직급", "성명", "생년월일", "종별", "사용기간(날짜)", "사용시간(시분)", "사유", "비고"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def _state(a_path, b_path, prev_path=None):
    state = AppState()
    state.config_obj = _config()
    state.load_files(str(a_path), str(b_path), prev_path)
    return state


def _july_output(tmp_path, monkeypatch):
    """7월: 경조사휴가 1건을 무급·미발생으로 확인하고 계산 -> 결과파일 경로를 돌려준다."""
    a, b = tmp_path / "a.xlsx", tmp_path / "b_july.xlsx"
    _write_a(a)
    _write_b(b, [["본부", "기간제", "홍길동", "1990-01-01", "경조사휴가", "2026-07-13", None, "", ""]])
    state = _state(a, b)
    state.confirm_special_leave([{"paid": False, "accrual": False}])
    state.batch_assign(list(state.people.keys()), "테스트조사")
    state.prepare_calculation(2026, 7)
    state.run_calculation()
    monkeypatch.setattr("webapp.state.downloads_dir", lambda: str(tmp_path))
    return a, state.download()[0]


def test_july_output_carries_status_and_decision(tmp_path, monkeypatch):
    _a, saved = _july_output(tmp_path, monkeypatch)
    wb = openpyxl.load_workbook(saved)
    assert STATUS_SHEET_NAME in wb.sheetnames
    ws = wb[STATUS_SHEET_NAME]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert pending.PAID_COL in headers and pending.ACCRUAL_COL in headers
    paid_col = headers.index(pending.PAID_COL) + 1
    accrual_col = headers.index(pending.ACCRUAL_COL) + 1
    assert ws.cell(row=2, column=paid_col).value == "무급"
    assert ws.cell(row=2, column=accrual_col).value == "미발생"


def test_august_only_b_file_passes_coverage_with_previous_file(tmp_path, monkeypatch):
    a, prev_path = _july_output(tmp_path, monkeypatch)

    b8 = tmp_path / "b_aug.xlsx"
    _write_b(b8, [["본부", "기간제", "홍길동", "1990-01-01", "연가", "2026-08-10", None, "", ""]])
    aug = _state(a, b8, prev_path)

    # 7월 건은 판정이 저장돼 있으므로 다시 묻지 않는다.
    assert aug.special_leave_groups() == []
    # 7월 근무상황이 전월 파일에서 채워져 계약 시작월 검증을 통과한다.
    aug.batch_assign(list(aug.people.keys()), "테스트조사")
    aug.prepare_calculation(2026, 8)
    assert [row["종별"] for row in aug.giganje_rows] == ["경조사휴가", "연가"]


def test_august_without_previous_file_is_blocked(tmp_path):
    a, b8 = tmp_path / "a.xlsx", tmp_path / "b_aug.xlsx"
    _write_a(a)
    _write_b(b8, [["본부", "기간제", "홍길동", "1990-01-01", "연가", "2026-08-10", None, "", ""]])
    state = _state(a, b8)
    state.batch_assign(list(state.people.keys()), "테스트조사")
    try:
        state.prepare_calculation(2026, 8)
    except ValueError as e:
        assert "전월" in str(e)
        return
    raise AssertionError("7월이 없으면 계산을 막아야 한다")
