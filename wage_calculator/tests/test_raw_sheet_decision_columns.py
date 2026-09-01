import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from core import pending
from core.parser import events_from_row
from output.raw_sheet import COLUMNS, SHEET_NAME, build_raw_status_sheet


def _row(category, period, **extra):
    row = {
        "소속": "본부", "직급": "기간제", "성명": "홍길동", "생년월일": "1990-01-01",
        "종별": category, "사용기간(날짜)": period, "사용시간(시분)": None,
        "사유": "", "비고": "", "연락처": "", "결재상태": "완료",
    }
    row.update(extra)
    return row


def test_columns_include_decisions():
    assert COLUMNS[-2:] == [pending.PAID_COL, pending.ACCRUAL_COL]


def test_sheet_writes_decision_values():
    row = _row("경조사휴가", "2026-08-10")
    row[pending.PAID_COL] = "무급"
    row[pending.ACCRUAL_COL] = "미발생"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_raw_status_sheet(wb, [row])
    ws = wb[SHEET_NAME]
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(COLUMNS) + 1)]
    assert headers[-2:] == [pending.PAID_COL, pending.ACCRUAL_COL]
    assert ws.cell(row=2, column=len(COLUMNS) - 1).value == "무급"
    assert ws.cell(row=2, column=len(COLUMNS)).value == "미발생"


def test_events_from_row_reads_decisions():
    row = _row("경조사휴가", "2026-08-10")
    row[pending.PAID_COL] = "무급"
    row[pending.ACCRUAL_COL] = "미발생"
    e = events_from_row(row)[0]
    assert e.decided is True
    assert e.unpaid is True
    assert e.breaks is True
    assert e.classified == "무급특별휴가"


def test_events_from_row_reads_paid_decision():
    row = _row("경조사휴가", "2026-08-10")
    row[pending.PAID_COL] = "유급"
    row[pending.ACCRUAL_COL] = "발생"
    e = events_from_row(row)[0]
    assert e.decided is True
    assert e.unpaid is False
    assert e.breaks is False
    assert e.classified == "유급특별휴가"


def test_known_category_decision_does_not_overwrite_classification():
    row = _row("연가", "2026-08-10", **{"사유": "집안일"})
    row[pending.PAID_COL] = "유급"
    row[pending.ACCRUAL_COL] = "발생"
    e = events_from_row(row)[0]
    assert e.classified == "연가"
    assert e.decided is True


def test_partial_decision_is_ignored():
    """한 칸만 채워진 행은 판정으로 인정하지 않고 다시 묻는다."""
    row = _row("경조사휴가", "2026-08-10")
    row[pending.PAID_COL] = "유급"
    e = events_from_row(row)[0]
    assert e.decided is False


def test_decided_row_is_not_collected_again():
    """전월 파일에서 판정을 이어받은 건은 확인 화면에 다시 올리지 않는다."""
    row = _row("경조사휴가", "2026-08-10")
    row[pending.PAID_COL] = "유급"
    row[pending.ACCRUAL_COL] = "발생"
    person = SimpleNamespace(name="홍길동", events=events_from_row(row))
    assert pending.collect_groups({"홍길동::1990-01-01": person}) == []


def test_undecided_row_is_still_collected():
    row = _row("경조사휴가", "2026-08-10")
    person = SimpleNamespace(name="홍길동", events=events_from_row(row))
    assert len(pending.collect_groups({"홍길동::1990-01-01": person})) == 1


if __name__ == "__main__":
    for _name, _fn in sorted(list(globals().items())):
        if _name.startswith("test_"):
            _fn()
    print("ALL OK")
