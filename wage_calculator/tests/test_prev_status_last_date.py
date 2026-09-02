import sys
from pathlib import Path
from datetime import date

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from core.parser import last_status_date
from output.raw_sheet import build_raw_status_sheet
from output.wage_sheet import build_wage_sheet
from webapp.state import AppState


def _row(category, period, time_field=None):
    return {
        "소속": "본부", "직급": "기간제", "성명": "홍길동", "생년월일": "1990-01-01",
        "종별": category, "사용기간(날짜)": period, "사용시간(시분)": time_field,
        "사유": "", "연락처": "", "결재상태": "완료", "비고": "",
    }


def test_returns_none_for_empty_rows():
    assert last_status_date([]) is None


def test_returns_latest_end_date():
    rows = [_row("연가", "2026-07-06"), _row("결근", "2026-07-20"), _row("공가", "2026-07-13")]
    assert last_status_date(rows) == date(2026, 7, 20)


def test_uses_end_of_range_not_start():
    rows = [_row("연가", "2026-07-06~2026-07-08")]
    assert last_status_date(rows) == date(2026, 7, 8)


def test_ignores_unreadable_dates():
    rows = [_row("연가", "2026-07-06"), _row("결근", "날짜아님"), _row("공가", None)]
    assert last_status_date(rows) == date(2026, 7, 6)


def _write_a(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["성명", "주민번호", "은행", "계좌번호", "생년월일"])
    ws.append(["홍길동", "900101-1234567", "농협", "123-456", "1990-01-01"])
    wb.save(path)


def _write_b(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["소속", "직급", "성명", "생년월일", "종별", "사용기간(날짜)", "사용시간(시분)", "사유", "비고"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def _write_prev(path, rows):
    """실제 산출물과 같은 모양(임금내역 시트 + 근무상황 시트)으로 만든다.
    근무상황 시트만 있으면 load_previous_payroll이 정상적으로 실패한다."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_wage_sheet(wb, [])
    build_raw_status_sheet(wb, rows)
    wb.save(path)


def test_upload_response_carries_last_date(tmp_path):
    a, b, prev = tmp_path / "a.xlsx", tmp_path / "b.xlsx", tmp_path / "prev.xlsx"
    _write_a(a)
    _write_b(b, [["본부", "기간제", "홍길동", "1990-01-01", "연가", "2026-08-10", None, "", ""]])
    _write_prev(prev, [_row("결근", "2026-07-06"), _row("공가", "2026-07-20")])

    state = AppState()
    result = state.load_files(str(a), str(b), str(prev))
    assert result["prev_status_last_date"] == "2026-07-20"


def test_upload_response_has_no_last_date_without_prev_file(tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _write_a(a)
    _write_b(b, [["본부", "기간제", "홍길동", "1990-01-01", "연가", "2026-08-10", None, "", ""]])

    state = AppState()
    result = state.load_files(str(a), str(b))
    assert result["prev_status_last_date"] is None


def test_reloading_clears_last_date(tmp_path):
    a, b, prev = tmp_path / "a.xlsx", tmp_path / "b.xlsx", tmp_path / "prev.xlsx"
    _write_a(a)
    _write_b(b, [["본부", "기간제", "홍길동", "1990-01-01", "연가", "2026-08-10", None, "", ""]])
    _write_prev(prev, [_row("결근", "2026-07-06")])

    state = AppState()
    state.load_files(str(a), str(b), str(prev))
    result = state.load_files(str(a), str(b))
    assert result["prev_status_last_date"] is None
    assert state.prev_status_last_date is None
