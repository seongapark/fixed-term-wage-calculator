"""한셀로 저장한 xlsx(셀 서식을 mc:AlternateContent로 감쌈)도 열려야 한다.

openpyxl은 AlternateContent를 건너뛰어 서식 목록이 짧아지고, 그 서식을 쓰는 셀에서
IndexError가 나 업로드가 500으로 실패했다.
"""
import re
import zipfile

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from core.parser import _open_workbook

MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def _hancell_like(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "성명"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "홍길동"
    ws["A2"].fill = PatternFill("solid", fgColor="FF4472C4")
    wb.save(path)

    with zipfile.ZipFile(path) as z:
        files = {i.filename: z.read(i.filename) for i in z.infolist()}
    styles = files["xl/styles.xml"].decode("utf-8")
    cell_xfs = re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles, re.S)
    xfs = re.findall(r"<xf\b.*?(?:/>|</xf>)", cell_xfs.group(1), re.S)
    # 첫 xf만 그대로 두고 나머지를 한셀처럼 Choice/Fallback으로 감싼다.
    wrapped = xfs[0] + "".join(
        f'<mc:AlternateContent {MC}><mc:Choice Requires="hs">{x}</mc:Choice>'
        f"<mc:Fallback>{x}</mc:Fallback></mc:AlternateContent>"
        for x in xfs[1:]
    )
    styles = styles[: cell_xfs.start(1)] + wrapped + styles[cell_xfs.end(1):]
    files["xl/styles.xml"] = styles.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


def test_hancell_saved_workbook_opens(tmp_path):
    path = tmp_path / "hancell.xlsx"
    _hancell_like(path)
    with pytest.raises(IndexError):
        openpyxl.load_workbook(path, data_only=True)  # 수정 전 동작 재현

    wb = _open_workbook(path)
    ws = wb.worksheets[0]
    assert ws["A1"].value == "성명"
    assert ws["A2"].value == "홍길동"
    assert ws["A1"].font.b
