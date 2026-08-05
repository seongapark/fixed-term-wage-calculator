import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from core.retroactive import DepartedRetro
from output.departed_sheet import build_departed_sheet, build_departed_evidence_sheet, COL


def test_departed_sheet_values():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    d = DepartedRetro(name="최도영", ssn="980126-2641395", bank="하나은행", account="111",
                       prev_recalculated=1950000, prev_paid=2000000, adjustment=-50000)
    ws = build_departed_sheet(wb, [d])

    assert ws.cell(row=2, column=COL["name"]).value == "최도영"
    assert ws.cell(row=2, column=COL["ssn"]).value == "980126-2641395"
    assert ws.cell(row=2, column=COL["prev_recalculated"]).value == 1950000
    assert ws.cell(row=2, column=COL["prev_paid"]).value == 2000000
    assert ws.cell(row=2, column=COL["adjustment"]).value == -50000
    print("OK: test_departed_sheet_values")


def test_departed_evidence_sheet_formula():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    d = DepartedRetro(name="최도영", ssn="980126-2641395", bank="하나은행", account="111",
                       prev_recalculated=1950000, prev_paid=2000000, adjustment=-50000)
    ws = build_departed_evidence_sheet(wb, [d])

    assert ws.cell(row=2, column=3).value == 1950000  # 전월재계산액
    assert ws.cell(row=2, column=4).value == 2000000  # 전월실지급액
    assert ws.cell(row=2, column=5).value == "=C2-D2"
    print("OK: test_departed_evidence_sheet_formula")


if __name__ == "__main__":
    test_departed_sheet_values()
    test_departed_evidence_sheet_formula()
    print("ALL OK")
